"""
PASSO 4 - Pipeline de integração (usina + inversores + meteorologia) e histórico anual

Este script tem duas partes:

1) processar_dia(...)   -> generaliza o integracao_usina.py original (que estava
   fixo em 22/09/2025) para qualquer data disponível na pasta de relatórios.
   A lógica de leitura/validação/junção é a mesma do script original (já testada
   e com correlação >0.9999 entre potência da usina e soma dos inversores);
   só parametrizei a data em vez de deixar fixa no código.

2) ler_historico_anual(...) -> lê os relatórios "estatístico (ALL)" (que ainda
   não tinham sido tratados por nenhum script) e resolve um problema real de
   formatação: os valores vêm como texto com pontos "duplicados"
   (ex: "367.169.80"), porque o separador decimal (vírgula) foi trocado por
   ponto na exportação, ficando ambíguo com o separador de milhar (que também
   é ponto). A regra usada: o ÚLTIMO ponto é sempre o decimal, os anteriores
   são separador de milhar.

Achados confirmados com este script (ver relatório Passo 4 para o detalhe):
- INV03 gerou 0.0 kWh entre 2020 e 2024 (5 anos) e só passou a gerar em 2025.
- A meteorologia tratada cobre 15/07/2022 a 31/07/2026 -> não cobre 22/09/2026
  ainda, então a integração dessa data fica sem irradiância/temperatura.
"""

import os
import re
import sys
import glob
from datetime import datetime
from functools import lru_cache

import numpy as np
import pandas as pd
from openpyxl import load_workbook

PASTA_PROJETO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(PASTA_PROJETO, "app"))
import catalogo_meteorologia as catalogo  # noqa: E402

PASTA_DADOS = os.path.join(PASTA_PROJETO, "dados")
PASTA_RELATORIOS = os.path.join(PASTA_DADOS, "brutos", "usina_powerstation")
PASTA_SAIDA = os.path.join(PASTA_DADOS, "processados")
ARQUIVO_METEOROLOGIA = os.path.join(
    PASTA_DADOS, "estudante1_meteorologia", "brutos", "dados_meteorologicos_tratados.xlsx"
)

POTENCIA_NOMINAL = {"INV01": 35.0, "INV02": 35.0, "INV03": 35.0, "INV04": 14.0}
POTENCIA_NOMINAL_USINA = 119.0

# ATENÇÃO: a potência nominal de 14 kW para o INV04 veio do script original
# (integracao_usina.py) e eu não confirmei essa informação em nenhuma
# especificação oficial. Os outros três são ~35 kW cada, batendo com o que
# já tínhamos mapeado no PowerStation. Isso fica como pendência a verificar.


# ============================================================
# PARTE 1 - INTEGRAÇÃO DIÁRIA (generalizada por data)
# ============================================================

def extrair_data_relatorio(valor):
    texto = str(valor)
    resultado = re.search(r"(\d{2}/\d{2}/\d{4})", texto)
    if resultado is None:
        raise ValueError("Não foi possível identificar a data do relatório.")
    return datetime.strptime(resultado.group(1), "%d/%m/%Y").date()


def conferir_data_no_arquivo(df, data_esperada, nome_arquivo):
    """Compara a data escrita dentro do relatório (linha 1, coluna 0) com a data
    do nome do arquivo. Se divergirem, avisa -- o nome pode ter sido renomeado
    ou o arquivo estar trocado."""
    try:
        data_interna = extrair_data_relatorio(df.iloc[1, 0])
    except ValueError:
        print(f"   [Aviso] Não achei a data dentro de {nome_arquivo}.")
        return
    if data_interna != data_esperada:
        print(f"   [Aviso] {nome_arquivo}: data interna {data_interna} difere da data do nome {data_esperada}.")


def criar_timestamps(data, horarios):
    timestamps = []
    for horario in horarios:
        if pd.isna(horario):
            timestamps.append(pd.NaT)
            continue
        texto = str(horario).strip()
        try:
            hora = datetime.strptime(texto, "%H:%M").time()
            timestamps.append(pd.Timestamp(datetime.combine(data, hora)))
        except ValueError:
            timestamps.append(pd.NaT)
    return timestamps


def arquivo_unico(padrao):
    """Devolve o arquivo que casa com o padrão glob. Falha com mensagem clara
    se não existir e avisa se houver mais de um (usa o mais recente pelo nome)."""
    achados = sorted(glob.glob(padrao))
    if not achados:
        raise FileNotFoundError(f"Nenhum arquivo encontrado para o padrão: {padrao}")
    if len(achados) > 1:
        print(f"   [Aviso] {len(achados)} arquivos casam com o padrão; usando {os.path.basename(achados[-1])}")
    return achados[-1]


def encontrar_datas_disponiveis():
    """Descobre, pelos nomes de arquivo, quais datas (YYYYMMDD) têm par
    completo (usina + inversores) disponível na pasta de relatórios."""
    padrao = re.compile(r"\((\d{8})\)")
    datas_usina = set()
    datas_inversor = set()
    for arq in glob.glob(os.path.join(PASTA_RELATORIOS, "Relatório de operação da usina(*)*.xlsx")):
        m = padrao.search(arq)
        if m:
            datas_usina.add(m.group(1))
    for arq in glob.glob(os.path.join(PASTA_RELATORIOS, "Relatório de operação do inversor(*)*.xlsx")):
        m = padrao.search(arq)
        if m:
            datas_inversor.add(m.group(1))
    return sorted(datas_usina & datas_inversor)


def ler_relatorio_usina(caminho, data_relatorio):
    df = pd.read_excel(caminho, header=None)
    conferir_data_no_arquivo(df, data_relatorio, os.path.basename(caminho))
    horarios = df.iloc[2, 2:].tolist()
    timestamps = criar_timestamps(data_relatorio, horarios)

    linha_potencia = None
    maior_quantidade = -1
    plantas_com_dado = 0
    for linha in range(len(df)):
        indicador = str(df.iloc[linha, 1]).strip()
        if indicador != "Potência FV(kW)":
            continue
        valores = pd.to_numeric(df.iloc[linha, 2:], errors="coerce")
        quantidade = int(valores.notna().sum())
        if quantidade > 0:
            plantas_com_dado += 1
        if quantidade > maior_quantidade:
            maior_quantidade = quantidade
            linha_potencia = linha

    if plantas_com_dado > 1:
        print("   [Aviso] Mais de uma planta tem 'Potência FV(kW)' com dado neste relatório; "
              "usando a linha com mais valores. Confira se é a planta certa.")

    if linha_potencia is None:
        raise ValueError("A linha 'Potência FV(kW)' não foi encontrada.")

    potencia = pd.to_numeric(df.iloc[linha_potencia, 2:], errors="coerce").tolist()
    resultado = pd.DataFrame({"timestamp": timestamps, "potencia_fv_kw": potencia})
    resultado = resultado.dropna(subset=["timestamp"]).sort_values("timestamp").reset_index(drop=True)
    return resultado


def extrair_inversor(texto):
    resultado = re.search(r"Nome do inversor:(\d+)", str(texto))
    if resultado is None:
        return None
    return f"INV{int(resultado.group(1)):02d}"


def ler_relatorio_inversores(caminho, data_relatorio):
    df = pd.read_excel(caminho, header=None)
    conferir_data_no_arquivo(df, data_relatorio, os.path.basename(caminho))
    horarios = df.iloc[2, 3:].tolist()
    timestamps = criar_timestamps(data_relatorio, horarios)

    resultados = []
    for linha in range(len(df)):
        indicador = str(df.iloc[linha, 2]).strip()
        if indicador != "Potência ativa(kW)":
            continue
        inversor = extrair_inversor(df.iloc[linha, 0])
        if inversor is None:
            continue
        valores = pd.to_numeric(df.iloc[linha, 3:], errors="coerce").tolist()
        resultados.append(pd.DataFrame({"timestamp": timestamps, inversor: valores}))

    if not resultados:
        raise ValueError("Nenhum registro de potência ativa foi encontrado.")

    resultado = resultados[0]
    for temp in resultados[1:]:
        resultado = resultado.merge(temp, on="timestamp", how="outer")
    resultado = resultado.sort_values("timestamp").reset_index(drop=True)

    for inversor in ["INV01", "INV02", "INV03", "INV04"]:
        if inversor not in resultado.columns:
            resultado[inversor] = np.nan

    return resultado[["timestamp", "INV01", "INV02", "INV03", "INV04"]]


def intervalo_amostragem_minutos(timestamps):
    """Intervalo típico (mediana) entre amostras, em minutos. Os relatórios de
    2025 vêm de 5 em 5 min e os de 2026 de 10 em 10 min."""
    ts = pd.Series(timestamps).dropna().sort_values()
    if len(ts) < 2:
        return None
    return float(ts.diff().dropna().median().total_seconds() / 60)


def preparar_dados_eletricos(usina, inversores):
    df = pd.merge(usina, inversores, on="timestamp", how="outer").sort_values("timestamp").reset_index(drop=True)
    colunas_inversores = ["INV01", "INV02", "INV03", "INV04"]

    df["inversores_com_dado"] = df[colunas_inversores].notna().sum(axis=1)
    df["quatro_inversores_com_dado"] = df[colunas_inversores].notna().all(axis=1)

    df["potencia_inversores_soma_kw"] = np.nan
    mascara_completa = df[colunas_inversores].notna().all(axis=1)
    df.loc[mascara_completa, "potencia_inversores_soma_kw"] = df.loc[mascara_completa, colunas_inversores].sum(axis=1)

    df["diferenca_potencia_kw"] = df["potencia_fv_kw"] - df["potencia_inversores_soma_kw"]
    df["diferenca_absoluta_kw"] = df["diferenca_potencia_kw"].abs()

    for inv in colunas_inversores:
        df[f"{inv}_relativa_%"] = 100 * df[inv] / POTENCIA_NOMINAL[inv]
    df["potencia_fv_relativa_%"] = 100 * df["potencia_fv_kw"] / POTENCIA_NOMINAL_USINA

    return df


@lru_cache(maxsize=1)
def ler_meteorologia_completa():
    """Lê a aba dados_tratados UMA vez por execução (o arquivo tem ~44 MB e
    abrir de novo a cada dia deixava o pipeline muito lento)."""
    return pd.read_excel(ARQUIVO_METEOROLOGIA, sheet_name="dados_tratados")


def carregar_meteorologia(data_relatorio):
    """Lê a aba dados_tratados da meteorologia e filtra pelo dia pedido.
    Retorna DataFrame vazio (com colunas) se o dia não estiver coberto."""
    df = ler_meteorologia_completa().copy()
    df["timestamp_meteorologia_original"] = pd.to_datetime(df["UTC -3"], errors="coerce")
    df = df.dropna(subset=["timestamp_meteorologia_original"])

    inicio = pd.Timestamp(data_relatorio)
    fim = inicio + pd.Timedelta(days=1)
    df = df[(df["timestamp_meteorologia_original"] >= inicio) & (df["timestamp_meteorologia_original"] < fim)]

    df = df.rename(columns={
        "Radiação solar global inst LPPYRA02_tratado": "ghi_lppyra02_w_m2",
        "Radiação solar global inst LPNET14_tratado": "ghi_lpnet14_w_m2",
        "Temperatura inst_tratado": "temperatura_c",
        "Velocidade vento med_tratado": "vento_m_s",
        "Precipitação_tratado": "precipitacao",
    })
    colunas = ["timestamp_meteorologia_original", "ghi_lppyra02_w_m2", "ghi_lpnet14_w_m2",
               "temperatura_c", "vento_m_s", "precipitacao"]
    colunas = [c for c in colunas if c in df.columns]
    return df[colunas].sort_values("timestamp_meteorologia_original").reset_index(drop=True)


def carregar_meteorologia_completa(data_relatorio):
    """Todas as variáveis do catálogo (51 valores + 24 flags) do dia pedido,
    prontas para meteorologia_15min. Os valores vêm da versão _tratado quando
    ela existe e do dado bruto quando não existe (ver o catálogo). Devolve
    DataFrame vazio, com as colunas certas, se o dia não estiver coberto."""
    df = ler_meteorologia_completa()
    necessarias = list(catalogo.mapa_valores_arquivo().values()) + \
        list(catalogo.mapa_flags_arquivo().values()) + ["UTC -3", "lacuna_temporal"]
    ausentes = [c for c in necessarias if c not in df.columns]
    if ausentes:
        raise ValueError(f"O arquivo meteorológico não tem as colunas esperadas: {ausentes}")

    ts = pd.to_datetime(df["UTC -3"], errors="coerce")
    inicio = pd.Timestamp(data_relatorio)
    mascara = ts.notna() & (ts >= inicio) & (ts < inicio + pd.Timedelta(days=1))
    dia = df.loc[mascara]

    saida = pd.DataFrame({
        "timestamp_15min": ts[mascara].dt.floor("15min"),
        "timestamp_original": ts[mascara],
        "lacuna_temporal": dia["lacuna_temporal"],
    })
    for coluna_banco, coluna_arquivo in catalogo.mapa_valores_arquivo().items():
        saida[coluna_banco] = dia[coluna_arquivo]
    for coluna_flag, coluna_arquivo in catalogo.mapa_flags_arquivo().items():
        saida[coluna_flag] = dia[coluna_arquivo]

    saida = saida[catalogo.colunas_meteorologia_15min()]
    return saida.sort_values("timestamp_15min").reset_index(drop=True)


def integrar_15min(df_eletrico, df_meteo):
    """Agrega o elétrico (5min) para 15min por média (mean) e só então
    cruza com a meteorologia, que já nasce em 15min. Cada timestamp_15min
    fica único (1 linha por bucket de 15min) -- é o que o banco espera."""
    df_eletrico = df_eletrico.copy()
    df_eletrico["timestamp_15min"] = df_eletrico["timestamp"].dt.floor("15min")

    colunas_numericas = [
        "potencia_fv_kw", "INV01", "INV02", "INV03", "INV04",
        "inversores_com_dado", "potencia_inversores_soma_kw",
        "diferenca_potencia_kw", "diferenca_absoluta_kw",
        "INV01_relativa_%", "INV02_relativa_%", "INV03_relativa_%", "INV04_relativa_%",
        "potencia_fv_relativa_%",
    ]
    colunas_numericas = [c for c in colunas_numericas if c in df_eletrico.columns]

    agregado = (
        df_eletrico
        .groupby("timestamp_15min")[colunas_numericas]
        .mean()
        .reset_index()
    )
    # quatro_inversores_com_dado: só é True no bucket de 15min se TODAS as
    # amostras de 5min daquele bucket tinham os 4 inversores com dado.
    agregado = agregado.merge(
        df_eletrico.groupby("timestamp_15min")["quatro_inversores_com_dado"].all().reset_index(),
        on="timestamp_15min",
    )

    if df_meteo.empty:
        df_completo = agregado.copy()
        for c in ["timestamp_meteorologia_original", "ghi_lppyra02_w_m2", "ghi_lpnet14_w_m2",
                  "temperatura_c", "vento_m_s", "precipitacao"]:
            df_completo[c] = pd.NA
        df_completo["meteorologia_disponivel"] = False
        df_completo["ghi_disponivel"] = False
        return df_completo

    df_meteo = df_meteo.copy()
    df_meteo["timestamp_15min"] = df_meteo["timestamp_meteorologia_original"].dt.floor("15min")

    df_completo = pd.merge(agregado, df_meteo, on="timestamp_15min", how="left")
    df_completo["meteorologia_disponivel"] = df_completo["timestamp_meteorologia_original"].notna()
    df_completo["ghi_disponivel"] = df_completo[["ghi_lppyra02_w_m2", "ghi_lpnet14_w_m2"]].notna().any(axis=1)
    return df_completo


def processar_dia(data_alvo_str):
    """data_alvo_str no formato YYYYMMDD. Retorna dict com os DataFrames prontos."""
    data_relatorio = datetime.strptime(data_alvo_str, "%Y%m%d").date()

    arq_usina = arquivo_unico(os.path.join(PASTA_RELATORIOS, f"Relatório de operação da usina({data_alvo_str})*.xlsx"))
    arq_inversor = arquivo_unico(os.path.join(PASTA_RELATORIOS, f"Relatório de operação do inversor({data_alvo_str})*.xlsx"))

    usina = ler_relatorio_usina(arq_usina, data_relatorio)
    inversores = ler_relatorio_inversores(arq_inversor, data_relatorio)
    eletrico = preparar_dados_eletricos(usina, inversores)

    intervalo = intervalo_amostragem_minutos(eletrico["timestamp"])
    if intervalo is not None and intervalo != 5:
        print(f"   [Aviso] Este relatório tem amostras a cada {intervalo:g} min (não 5). "
              "O nome 'registros_eletricos_5min' no resumo conta amostras, não 5 minutos.")

    meteo = carregar_meteorologia(data_relatorio)
    integracao_15min = integrar_15min(eletrico, meteo)
    meteorologia_completa = carregar_meteorologia_completa(data_relatorio)

    comparaveis = eletrico.dropna(subset=["potencia_fv_kw", "potencia_inversores_soma_kw"])
    correlacao = comparaveis["potencia_fv_kw"].corr(comparaveis["potencia_inversores_soma_kw"]) if len(comparaveis) > 1 else None

    resumo = {
        "data": data_relatorio.isoformat(),
        "registros_eletricos_5min": len(eletrico),
        "registros_com_potencia_usina": int(eletrico["potencia_fv_kw"].notna().sum()),
        "registros_com_4_inversores": int(eletrico["quatro_inversores_com_dado"].sum()),
        "registros_15min": len(integracao_15min),
        "registros_com_meteorologia": int(integracao_15min["meteorologia_disponivel"].sum()) if "meteorologia_disponivel" in integracao_15min else 0,
        "correlacao_usina_vs_soma_inversores": round(correlacao, 8) if correlacao is not None else None,
        "geracao_max_inv01_kw": round(inversores["INV01"].max(), 2) if inversores["INV01"].notna().any() else None,
        "geracao_max_inv02_kw": round(inversores["INV02"].max(), 2) if inversores["INV02"].notna().any() else None,
        "geracao_max_inv03_kw": round(inversores["INV03"].max(), 2) if inversores["INV03"].notna().any() else None,
        "geracao_max_inv04_kw": round(inversores["INV04"].max(), 2) if inversores["INV04"].notna().any() else None,
    }

    return {"eletrico_5min": eletrico, "integracao_15min": integracao_15min,
            "meteorologia_15min": meteorologia_completa, "resumo": resumo}


def salvar_dia(resultado, data_alvo_str):
    os.makedirs(PASTA_SAIDA, exist_ok=True)
    caminho_saida = os.path.join(PASTA_SAIDA, f"integracao_usina_{data_alvo_str}.xlsx")
    with pd.ExcelWriter(caminho_saida, engine="openpyxl") as writer:
        resultado["eletrico_5min"].to_excel(writer, sheet_name="eletrico_5min", index=False)
        resultado["integracao_15min"].to_excel(writer, sheet_name="integracao_15min", index=False)
        resultado["meteorologia_15min"].to_excel(writer, sheet_name="meteorologia_15min", index=False)
        pd.DataFrame([resultado["resumo"]]).T.reset_index().rename(
            columns={"index": "indicador", 0: "valor"}
        ).to_excel(writer, sheet_name="resumo", index=False)
    return caminho_saida


# ============================================================
# PARTE 2 - HISTÓRICO ANUAL (relatórios "estatístico ALL")
# ============================================================

def parse_numero_ptbr_ambiguo(valor):
    """Converte textos como '367.169.80' ou '27.795.60' em float, assumindo
    que o ÚLTIMO ponto é o separador decimal e os anteriores são de milhar."""
    if valor is None:
        return None
    texto = str(valor).strip()
    if texto in ("", "--", "NA", "None"):
        return None
    partes = texto.split(".")
    if len(partes) == 1:
        try:
            return float(texto)
        except ValueError:
            return None
    inteiro = "".join(partes[:-1])
    decimal = partes[-1]
    try:
        return float(f"{inteiro}.{decimal}")
    except ValueError:
        return None


def ler_historico_anual_inversores(caminho):
    wb = load_workbook(caminho, data_only=True)
    ws = wb["Sheet0"]
    header = [c.value for c in ws[4]]
    colunas_anos = header[2:]  # "Valor total", "2020", ..., "2026"

    registros = []
    inversor_atual = None
    for row in ws.iter_rows(min_row=5, values_only=True):
        info, indicador, *valores = row
        if info:
            m = re.search(r"Nome do inversor:(\d+)", str(info))
            if m:
                inversor_atual = f"INV{int(m.group(1)):02d}"
        if indicador == "Geração de energia(kWh)":
            vals = [parse_numero_ptbr_ambiguo(v) for v in valores]
            registros.append([inversor_atual] + vals)

    return pd.DataFrame(registros, columns=["inversor"] + colunas_anos)


def ler_historico_anual_usina(caminho):
    wb = load_workbook(caminho, data_only=True)
    ws = wb["Sheet0"]
    header = [c.value for c in ws[4]]
    colunas_anos = header[2:]

    for row in ws.iter_rows(min_row=5, values_only=True):
        info, indicador, *valores = row
        if indicador == "Geração de energia(kWh)":
            vals = [parse_numero_ptbr_ambiguo(v) for v in valores]
            return pd.DataFrame([["USINA"] + vals], columns=["fonte"] + colunas_anos)
    return None


def validar_historico_anual(df_inversores, df_usina, tolerancia_kwh=1.0):
    """Confere se, em cada ano, a soma dos 4 inversores bate com a geração da
    usina. Devolve a lista de (coluna, diferença em kWh) que passou da tolerância."""
    divergencias = []
    if df_usina is None or df_inversores.empty:
        return divergencias
    for coluna in df_usina.columns[1:]:
        soma = df_inversores[coluna].sum()
        usina = df_usina[coluna].iloc[0]
        if pd.notna(soma) and pd.notna(usina) and abs(soma - usina) > tolerancia_kwh:
            divergencias.append((coluna, round(float(soma - usina), 2)))
    return divergencias


def cobertura_meteorologia():
    ts = pd.to_datetime(ler_meteorologia_completa()["UTC -3"], errors="coerce").dropna()
    return ts.min(), ts.max()


# ============================================================
# EXECUÇÃO
# ============================================================

if __name__ == "__main__":

    print("=" * 70)
    print("PASSO 4 - PIPELINE GENERALIZADA")
    print("=" * 70)

    datas = encontrar_datas_disponiveis()
    print(f"\nDatas com relatório completo (usina + inversores): {datas}")

    resumos = []
    for data_alvo in datas:
        print(f"\n--- Processando {data_alvo} ---")
        resultado = processar_dia(data_alvo)
        caminho = salvar_dia(resultado, data_alvo)
        print(f"Salvo em: {caminho}")
        print(resultado["resumo"])
        resumos.append(resultado["resumo"])

    print("\n" + "=" * 70)
    print("HISTÓRICO ANUAL (relatórios estatístico ALL)")
    print("=" * 70)

    arq_hist_inv = arquivo_unico(os.path.join(PASTA_RELATORIOS, "Relatório estatístico do inversor(ALL)*.xlsx"))
    arq_hist_usina = arquivo_unico(os.path.join(PASTA_RELATORIOS, "Relatório estatístico da usina(ALL)*.xlsx"))

    df_hist_inv = ler_historico_anual_inversores(arq_hist_inv)
    df_hist_usina = ler_historico_anual_usina(arq_hist_usina)

    print("\nGeração anual por inversor (kWh):")
    print(df_hist_inv.to_string(index=False))

    print("\nGeração anual da usina (kWh):")
    print(df_hist_usina.to_string(index=False))

    divergencias = validar_historico_anual(df_hist_inv, df_hist_usina)
    if divergencias:
        print(f"\n[Aviso] Soma dos inversores difere da usina em: {divergencias}")
    else:
        print("\nValidação: a soma dos 4 inversores bate com a usina em todos os anos.")

    caminho_hist = os.path.join(PASTA_SAIDA, "historico_anual_inversores.xlsx")
    with pd.ExcelWriter(caminho_hist, engine="openpyxl") as writer:
        df_hist_inv.to_excel(writer, sheet_name="por_inversor", index=False)
        df_hist_usina.to_excel(writer, sheet_name="usina_total", index=False)
    print(f"\nSalvo em: {caminho_hist}")

    inicio_meteo, fim_meteo = cobertura_meteorologia()
    print(f"\nCobertura temporal da meteorologia tratada: {inicio_meteo} até {fim_meteo}")

    print("\n" + "=" * 70)
    print("PASSO 4 CONCLUÍDO")
    print("=" * 70)
