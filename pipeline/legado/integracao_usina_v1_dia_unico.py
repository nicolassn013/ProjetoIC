import os
import re
import glob
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment


# ============================================================
# CONFIGURAÇÕES
# ============================================================

PASTA_PROJETO = os.path.dirname(
    os.path.abspath(__file__)
)

PASTA_DADOS = os.path.join(
    PASTA_PROJETO,
    "dados"
)

# Pasta onde estão os relatórios da usina/inversores
PASTA_RELATORIOS = os.path.join(
    PASTA_DADOS,
    "Relátorio da usina_inversores"
)

# ============================================================
# ARQUIVO METEOROLÓGICO
# ============================================================

ARQUIVO_METEOROLOGIA = os.path.join(
    PASTA_DADOS,
    "dados_meteorologicos_tratados.xlsx"
)

# ============================================================
# PROCURAR AUTOMATICAMENTE OS RELATÓRIOS DE 22/09/2025
# ============================================================

ARQUIVOS_USINA = glob.glob(
    os.path.join(
        PASTA_RELATORIOS,
        "Relatório de operação da usina(20250922)*.xlsx"
    )
)

ARQUIVOS_INVERSORES = glob.glob(
    os.path.join(
        PASTA_RELATORIOS,
        "Relatório de operação do inversor(20250922)*.xlsx"
    )
)

if not ARQUIVOS_USINA:

    raise FileNotFoundError(
        "\nNão foi encontrado o relatório de operação "
        "da usina de 22/09/2025."
        f"\nPasta pesquisada:\n{PASTA_RELATORIOS}"
    )

if not ARQUIVOS_INVERSORES:

    raise FileNotFoundError(
        "\nNão foi encontrado o relatório de operação "
        "dos inversores de 22/09/2025."
        f"\nPasta pesquisada:\n{PASTA_RELATORIOS}"
    )

# Se houver mais de um arquivo correspondente,
# usamos o primeiro encontrado.
ARQUIVO_USINA = ARQUIVOS_USINA[0]

ARQUIVO_INVERSORES = ARQUIVOS_INVERSORES[0]

# ============================================================
# ARQUIVO DE SAÍDA
# ============================================================

ARQUIVO_SAIDA = os.path.join(
    PASTA_DADOS,
    "integracao_usina_20250922.xlsx"
)

# ============================================================
# POTÊNCIA NOMINAL DOS INVERSORES
# ============================================================

POTENCIA_NOMINAL = {
    "INV01": 35.0,
    "INV02": 35.0,
    "INV03": 35.0,
    "INV04": 14.0
}

POTENCIA_NOMINAL_USINA = 119.0


# ============================================================
# LIMPAR TERMINAL
# ============================================================

def limpar_terminal():

    if os.name == "nt":
        os.system("cls")
    else:
        os.system("clear")


# ============================================================
# VERIFICAR ARQUIVOS
# ============================================================

def verificar_arquivos():

    arquivos = [
        ARQUIVO_METEOROLOGIA,
        ARQUIVO_USINA,
        ARQUIVO_INVERSORES
    ]

    for arquivo in arquivos:

        if not os.path.exists(arquivo):

            raise FileNotFoundError(
                f"\nArquivo não encontrado:\n{arquivo}"
            )

    print("\nArquivos encontrados:")

    print(
        f"\nMeteorologia:\n{ARQUIVO_METEOROLOGIA}"
    )

    print(
        f"\nUsina:\n{ARQUIVO_USINA}"
    )

    print(
        f"\nInversores:\n{ARQUIVO_INVERSORES}"
    )


# ============================================================
# EXTRAIR DATA DO RELATÓRIO
# ============================================================

def extrair_data_relatorio(valor):

    texto = str(valor)

    resultado = re.search(
        r"(\d{2}/\d{2}/\d{4})",
        texto
    )

    if resultado is None:

        raise ValueError(
            "Não foi possível identificar a data do relatório."
        )

    return datetime.strptime(
        resultado.group(1),
        "%d/%m/%Y"
    ).date()


# ============================================================
# CRIAR TIMESTAMPS
# ============================================================

def criar_timestamps(
    data,
    horarios
):

    timestamps = []

    for horario in horarios:

        if pd.isna(horario):

            timestamps.append(
                pd.NaT
            )

            continue

        texto = str(
            horario
        ).strip()

        try:

            hora = datetime.strptime(
                texto,
                "%H:%M"
            ).time()

            timestamp = pd.Timestamp(
                datetime.combine(
                    data,
                    hora
                )
            )

            timestamps.append(
                timestamp
            )

        except ValueError:

            timestamps.append(
                pd.NaT
            )

    return timestamps


# ============================================================
# LER RELATÓRIO DA USINA
# ============================================================

def ler_relatorio_usina():

    print(
        "\n" + "=" * 70
    )

    print(
        "LENDO RELATÓRIO DA USINA"
    )

    print(
        "=" * 70
    )

    df = pd.read_excel(
        ARQUIVO_USINA,
        header=None
    )

    # --------------------------------------------------------
    # Data do relatório
    # --------------------------------------------------------

    data_relatorio = (
        extrair_data_relatorio(
            df.iloc[1, 0]
        )
    )

    print(
        f"Data do relatório: "
        f"{data_relatorio}"
    )

    # --------------------------------------------------------
    # Horários
    #
    # Na planilha da usina:
    # coluna 1 = indicador
    # coluna 2 em diante = horários
    # --------------------------------------------------------

    horarios = df.iloc[
        2,
        2:
    ].tolist()

    timestamps = criar_timestamps(
        data_relatorio,
        horarios
    )

    # --------------------------------------------------------
    # Localizar a linha da Potência FV
    #
    # Existe mais de uma ocorrência no arquivo.
    # Selecionamos aquela com maior quantidade de valores.
    # --------------------------------------------------------

    linha_potencia = None

    maior_quantidade = -1

    for linha in range(
        len(df)
    ):

        indicador = str(
            df.iloc[
                linha,
                1
            ]
        ).strip()

        if indicador != "Potência FV(kW)":
            continue

        valores = pd.to_numeric(
            df.iloc[
                linha,
                2:
            ],
            errors="coerce"
        )

        quantidade = int(
            valores.notna().sum()
        )

        if quantidade > maior_quantidade:

            maior_quantidade = quantidade
            linha_potencia = linha

    if linha_potencia is None:

        raise ValueError(
            "A linha 'Potência FV(kW)' não foi encontrada."
        )

    potencia = pd.to_numeric(
        df.iloc[
            linha_potencia,
            2:
        ],
        errors="coerce"
    ).tolist()

    resultado = pd.DataFrame(
        {
            "timestamp": timestamps,
            "potencia_fv_kw": potencia
        }
    )

    resultado = resultado.dropna(
        subset=["timestamp"]
    )

    resultado = resultado.sort_values(
        "timestamp"
    ).reset_index(
        drop=True
    )

    print(
        f"Registros horários: "
        f"{len(resultado)}"
    )

    print(
        f"Registros com potência FV: "
        f"{resultado['potencia_fv_kw'].notna().sum()}"
    )

    print(
        f"Potência FV máxima: "
        f"{resultado['potencia_fv_kw'].max():.2f} kW"
    )

    return resultado


# ============================================================
# EXTRAIR NÚMERO DO INVERSOR
# ============================================================

def extrair_inversor(texto):

    resultado = re.search(
        r"Nome do inversor:(\d+)",
        str(texto)
    )

    if resultado is None:
        return None

    numero = int(
        resultado.group(1)
    )

    return f"INV{numero:02d}"


# ============================================================
# LER RELATÓRIO DOS INVERSORES
# ============================================================

def ler_relatorio_inversores():

    print(
        "\n" + "=" * 70
    )

    print(
        "LENDO RELATÓRIO DOS INVERSORES"
    )

    print(
        "=" * 70
    )

    df = pd.read_excel(
        ARQUIVO_INVERSORES,
        header=None
    )

    # --------------------------------------------------------
    # Data
    # --------------------------------------------------------

    data_relatorio = (
        extrair_data_relatorio(
            df.iloc[1, 0]
        )
    )

    print(
        f"Data do relatório: "
        f"{data_relatorio}"
    )

    # --------------------------------------------------------
    # Horários
    #
    # coluna 2 = indicador
    # coluna 3 em diante = horários
    # --------------------------------------------------------

    horarios = df.iloc[
        2,
        3:
    ].tolist()

    timestamps = criar_timestamps(
        data_relatorio,
        horarios
    )

    resultados = []

    # --------------------------------------------------------
    # Localizar cada linha de potência ativa
    # --------------------------------------------------------

    for linha in range(
        len(df)
    ):

        indicador = str(
            df.iloc[
                linha,
                2
            ]
        ).strip()

        if indicador != "Potência ativa(kW)":
            continue

        inversor = extrair_inversor(
            df.iloc[
                linha,
                0
            ]
        )

        if inversor is None:
            continue

        valores = pd.to_numeric(
            df.iloc[
                linha,
                3:
            ],
            errors="coerce"
        ).tolist()

        temp = pd.DataFrame(
            {
                "timestamp": timestamps,
                inversor: valores
            }
        )

        resultados.append(
            temp
        )

    if not resultados:

        raise ValueError(
            "Nenhum registro de potência ativa "
            "foi encontrado."
        )

    # --------------------------------------------------------
    # Juntar os quatro inversores
    # --------------------------------------------------------

    resultado = resultados[0]

    for temp in resultados[1:]:

        resultado = resultado.merge(
            temp,
            on="timestamp",
            how="outer"
        )

    resultado = resultado.sort_values(
        "timestamp"
    ).reset_index(
        drop=True
    )

    # --------------------------------------------------------
    # Garantir que os quatro inversores existam
    # --------------------------------------------------------

    for inversor in [
        "INV01",
        "INV02",
        "INV03",
        "INV04"
    ]:

        if inversor not in resultado.columns:

            resultado[inversor] = pd.NA

    resultado = resultado[
        [
            "timestamp",
            "INV01",
            "INV02",
            "INV03",
            "INV04"
        ]
    ]

    print(
        "\nRegistros por inversor:"
    )

    for inversor in [
        "INV01",
        "INV02",
        "INV03",
        "INV04"
    ]:

        quantidade = int(
            resultado[
                inversor
            ]
            .notna()
            .sum()
        )

        potencia_maxima = (
            resultado[
                inversor
            ]
            .max()
        )

        print(
            f"{inversor}: "
            f"{quantidade} registros | "
            f"máximo = "
            f"{potencia_maxima:.2f} kW"
        )

    return resultado


# ============================================================
# PREPARAR DADOS ELÉTRICOS
# ============================================================

def preparar_dados_eletricos(
    usina,
    inversores
):

    print(
        "\nPreparando dados elétricos..."
    )

    df = pd.merge(
        usina,
        inversores,
        on="timestamp",
        how="outer"
    )

    df = df.sort_values(
        "timestamp"
    ).reset_index(
        drop=True
    )

    colunas_inversores = [
        "INV01",
        "INV02",
        "INV03",
        "INV04"
    ]

    # --------------------------------------------------------
    # Quantidade de inversores com registro
    # --------------------------------------------------------

    df[
        "inversores_com_dado"
    ] = (
        df[
            colunas_inversores
        ]
        .notna()
        .sum(
            axis=1
        )
    )

    # --------------------------------------------------------
    # Todos os quatro inversores possuem dados?
    # --------------------------------------------------------

    df[
        "quatro_inversores_com_dado"
    ] = (
        df[
            colunas_inversores
        ]
        .notna()
        .all(
            axis=1
        )
    )

    # --------------------------------------------------------
    # Soma dos inversores
    #
    # SOMENTE quando os quatro possuem dados.
    #
    # Ausência de comunicação NÃO vira 0.
    # --------------------------------------------------------

    df[
        "potencia_inversores_soma_kw"
    ] = pd.NA

    mascara_completa = (
        df[
            colunas_inversores
        ]
        .notna()
        .all(
            axis=1
        )
    )

    df.loc[
        mascara_completa,
        "potencia_inversores_soma_kw"
    ] = (
        df.loc[
            mascara_completa,
            colunas_inversores
        ]
        .sum(
            axis=1
        )
    )

    # --------------------------------------------------------
    # Diferença entre a potência reportada pela usina
    # e a soma dos inversores.
    # --------------------------------------------------------

    df[
        "diferenca_potencia_kw"
    ] = (
        df[
            "potencia_fv_kw"
        ]
        - df[
            "potencia_inversores_soma_kw"
        ]
    )

    df[
        "diferenca_absoluta_kw"
    ] = (
        df[
            "diferenca_potencia_kw"
        ]
        .abs()
    )

    # --------------------------------------------------------
    # Potência relativa dos inversores
    # --------------------------------------------------------

    for inversor, potencia_nominal in (
        POTENCIA_NOMINAL.items()
    ):

        df[
            f"{inversor}_relativa_%"
        ] = (
            pd.to_numeric(
                df[inversor],
                errors="coerce"
            )
            / potencia_nominal
            * 100
        )

    # --------------------------------------------------------
    # Potência relativa da usina
    # --------------------------------------------------------

    df[
        "potencia_fv_relativa_%"
    ] = (
        pd.to_numeric(
            df[
                "potencia_fv_kw"
            ],
            errors="coerce"
        )
        / POTENCIA_NOMINAL_USINA
        * 100
    )

    return df


# ============================================================
# ESCOLHER COLUNA TRATADA DA METEOROLOGIA
# ============================================================

def obter_coluna_meteorologica(
    df,
    coluna_original
):

    coluna_tratada = (
        f"{coluna_original}_tratado"
    )

    if coluna_tratada in df.columns:

        return coluna_tratada

    return coluna_original


# ============================================================
# PREPARAR METEOROLOGIA
# ============================================================

def preparar_meteorologia():

    print(
        "\n" + "=" * 70
    )

    print(
        "CARREGANDO METEOROLOGIA"
    )

    print(
        "=" * 70
    )

    df_original = pd.read_excel(
        ARQUIVO_METEOROLOGIA,
        sheet_name="dados_tratados"
    )

    # --------------------------------------------------------
    # Colunas necessárias
    # --------------------------------------------------------

    colunas_necessarias = [
        "UTC -3",
        "Radiação solar global inst LPPYRA02",
        "Radiação solar global inst LPNET14",
        "Temperatura inst",
        "Velocidade vento med",
        "Precipitação"
    ]

    for coluna in colunas_necessarias:

        if coluna not in df_original.columns:

            raise ValueError(
                f"A coluna '{coluna}' "
                "não foi encontrada no arquivo meteorológico."
            )

    # --------------------------------------------------------
    # Selecionar original ou tratado
    # --------------------------------------------------------

    coluna_ghi_1 = obter_coluna_meteorologica(
        df_original,
        "Radiação solar global inst LPPYRA02"
    )

    coluna_ghi_2 = obter_coluna_meteorologica(
        df_original,
        "Radiação solar global inst LPNET14"
    )

    coluna_temperatura = obter_coluna_meteorologica(
        df_original,
        "Temperatura inst"
    )

    coluna_vento = obter_coluna_meteorologica(
        df_original,
        "Velocidade vento med"
    )

    coluna_precipitacao = obter_coluna_meteorologica(
        df_original,
        "Precipitação"
    )

    # --------------------------------------------------------
    # Criar DataFrame
    # --------------------------------------------------------

    df = df_original[
        [
            "UTC -3",
            coluna_ghi_1,
            coluna_ghi_2,
            coluna_temperatura,
            coluna_vento,
            coluna_precipitacao
        ]
    ].copy()

    df = df.rename(
        columns={
            "UTC -3":
                "timestamp_meteorologia_original",

            coluna_ghi_1:
                "ghi_lppyra02_w_m2",

            coluna_ghi_2:
                "ghi_lpnet14_w_m2",

            coluna_temperatura:
                "temperatura_c",

            coluna_vento:
                "vento_m_s",

            coluna_precipitacao:
                "precipitacao"
        }
    )

    # --------------------------------------------------------
    # Converter timestamp
    # --------------------------------------------------------

    df[
        "timestamp_meteorologia_original"
    ] = pd.to_datetime(
        df[
            "timestamp_meteorologia_original"
        ],
        errors="coerce"
    )

    df = df.dropna(
        subset=[
            "timestamp_meteorologia_original"
        ]
    )

    # --------------------------------------------------------
    # Selecionar somente 22/09/2025
    # --------------------------------------------------------

    inicio = pd.Timestamp(
        "2025-09-22 00:00:00"
    )

    fim = pd.Timestamp(
        "2025-09-23 00:00:00"
    )

    df = df[
        (
            df[
                "timestamp_meteorologia_original"
            ] >= inicio
        )
        &
        (
            df[
                "timestamp_meteorologia_original"
            ] < fim
        )
    ].copy()

    # --------------------------------------------------------
    # Normalizar para horários de 15 minutos.
    #
    # Exemplo:
    # 10:15:07 -> 10:15:00
    # --------------------------------------------------------

    df[
        "timestamp_15min"
    ] = (
        df[
            "timestamp_meteorologia_original"
        ]
        .dt.round(
            "15min"
        )
    )

    # --------------------------------------------------------
    # Verificar distância entre o horário real e o horário
    # de 15 minutos.
    # --------------------------------------------------------

    diferenca_segundos = (
        (
            df[
                "timestamp_meteorologia_original"
            ]
            -
            df[
                "timestamp_15min"
            ]
        )
        .dt.total_seconds()
        .abs()
    )

    # Tolerância de 2 minutos.
    df = df[
        diferenca_segundos <= 120
    ].copy()

    # --------------------------------------------------------
    # Evitar dois registros para o mesmo horário de 15 min.
    # --------------------------------------------------------

    df = df.drop_duplicates(
        subset=[
            "timestamp_15min"
        ],
        keep="first"
    )

    df = df.sort_values(
        "timestamp_15min"
    ).reset_index(
        drop=True
    )

    print(
        f"Registros meteorológicos no dia: "
        f"{len(df)}"
    )

    return df


# ============================================================
# INTEGRAR EM 15 MINUTOS
# ============================================================

def integrar_15_minutos(
    eletricos,
    meteorologia
):

    print(
        "\n" + "=" * 70
    )

    print(
        "CRIANDO INTEGRAÇÃO EM 15 MINUTOS"
    )

    print(
        "=" * 70
    )

    eletricos = eletricos.copy()

    # --------------------------------------------------------
    # Selecionar os horários de 15 em 15 minutos
    # a partir dos dados de 5 minutos.
    #
    # Usamos o timestamp real.
    # Não usamos posição de linha.
    # --------------------------------------------------------

    eletricos[
        "minuto"
    ] = (
        eletricos[
            "timestamp"
        ]
        .dt.minute
    )

    eletricos_15 = eletricos[
        eletricos[
            "minuto"
        ].isin(
            [
                0,
                15,
                30,
                45
            ]
        )
    ].copy()

    eletricos_15 = eletricos_15.drop(
        columns=[
            "minuto"
        ]
    )

    # --------------------------------------------------------
    # Criar timestamp de integração
    # --------------------------------------------------------

    eletricos_15[
        "timestamp_15min"
    ] = (
        eletricos_15[
            "timestamp"
        ]
        .dt.floor(
            "15min"
        )
    )

    # --------------------------------------------------------
    # Cruzar com meteorologia
    # --------------------------------------------------------

    resultado = pd.merge(
        eletricos_15,
        meteorologia,
        on="timestamp_15min",
        how="left"
    )

    # --------------------------------------------------------
    # Indicadores de disponibilidade
    # --------------------------------------------------------

    resultado[
        "meteorologia_disponivel"
    ] = (
        resultado[
            [
                "ghi_lppyra02_w_m2",
                "ghi_lpnet14_w_m2",
                "temperatura_c",
                "vento_m_s"
            ]
        ]
        .notna()
        .any(
            axis=1
        )
    )

    resultado[
        "ghi_disponivel"
    ] = (
        resultado[
            "ghi_lppyra02_w_m2"
        ].notna()
        |
        resultado[
            "ghi_lpnet14_w_m2"
        ].notna()
    )

    resultado = resultado.sort_values(
        "timestamp_15min"
    ).reset_index(
        drop=True
    )

    return resultado


# ============================================================
# VALIDAR POTÊNCIA DA USINA × INVERSORES
# ============================================================

def validar_potencia(
    eletricos
):

    print(
        "\n" + "=" * 70
    )

    print(
        "VALIDANDO POTÊNCIA DA USINA × INVERSORES"
    )

    print(
        "=" * 70
    )

    mascara = (
        eletricos[
            "potencia_fv_kw"
        ].notna()
        &
        eletricos[
            "potencia_inversores_soma_kw"
        ].notna()
    )

    validacao = eletricos.loc[
        mascara,
        [
            "timestamp",
            "potencia_fv_kw",
            "potencia_inversores_soma_kw",
            "diferenca_potencia_kw",
            "diferenca_absoluta_kw"
        ]
    ].copy()

    if validacao.empty:

        resumo = pd.DataFrame(
            {
                "indicador": [
                    "Registros comparáveis"
                ],
                "valor": [
                    0
                ]
            }
        )

        return validacao, resumo

    quantidade = len(
        validacao
    )

    diferenca_media = (
        validacao[
            "diferenca_potencia_kw"
        ]
        .mean()
    )

    erro_absoluto_medio = (
        validacao[
            "diferenca_absoluta_kw"
        ]
        .mean()
    )

    diferenca_maxima = (
        validacao[
            "diferenca_absoluta_kw"
        ]
        .max()
    )

    correlacao = (
        validacao[
            "potencia_fv_kw"
        ]
        .corr(
            validacao[
                "potencia_inversores_soma_kw"
            ]
        )
    )

    resumo = pd.DataFrame(
        {
            "indicador": [
                "Registros comparáveis",
                "Diferença média (kW)",
                "Erro absoluto médio (kW)",
                "Maior diferença absoluta (kW)",
                "Correlação"
            ],
            "valor": [
                quantidade,
                diferenca_media,
                erro_absoluto_medio,
                diferenca_maxima,
                correlacao
            ]
        }
    )

    print(
        f"Registros comparáveis: "
        f"{quantidade}"
    )

    print(
        f"Diferença média: "
        f"{diferenca_media:.4f} kW"
    )

    print(
        f"Erro absoluto médio: "
        f"{erro_absoluto_medio:.4f} kW"
    )

    print(
        f"Maior diferença absoluta: "
        f"{diferenca_maxima:.4f} kW"
    )

    print(
        f"Correlação: "
        f"{correlacao:.8f}"
    )

    return validacao, resumo


# ============================================================
# GERAR RESUMO
# ============================================================

def gerar_resumo(
    eletricos,
    integracao_15,
    validacao_resumo
):

    print(
        "\nGerando resumo..."
    )

    registros_5min = len(
        eletricos
    )

    potencia_usina = int(
        eletricos[
            "potencia_fv_kw"
        ]
        .notna()
        .sum()
    )

    registros_4_inversores = int(
        eletricos[
            "quatro_inversores_com_dado"
        ]
        .sum()
    )

    registros_meteorologia = int(
        integracao_15[
            "meteorologia_disponivel"
        ]
        .sum()
    )

    registros_ghi = int(
        integracao_15[
            "ghi_disponivel"
        ]
        .sum()
    )

    resumo = pd.DataFrame(
        {
            "indicador": [
                "Data analisada",
                "Registros elétricos de 5 min",
                "Registros com potência da usina",
                "Registros com os 4 inversores",
                "Registros integrados de 15 min",
                "Registros de 15 min com meteorologia",
                "Registros de 15 min com GHI",
                "Potência nominal da usina (kW)",
                "Potência nominal INV01 (kW)",
                "Potência nominal INV02 (kW)",
                "Potência nominal INV03 (kW)",
                "Potência nominal INV04 (kW)"
            ],
            "valor": [
                "22/09/2025",
                registros_5min,
                potencia_usina,
                registros_4_inversores,
                len(integracao_15),
                registros_meteorologia,
                registros_ghi,
                POTENCIA_NOMINAL_USINA,
                POTENCIA_NOMINAL["INV01"],
                POTENCIA_NOMINAL["INV02"],
                POTENCIA_NOMINAL["INV03"],
                POTENCIA_NOMINAL["INV04"]
            ]
        }
    )

    if not validacao_resumo.empty:

        resumo = pd.concat(
            [
                resumo,
                validacao_resumo
            ],
            ignore_index=True
        )

    return resumo


# ============================================================
# FORMATAR EXCEL
# ============================================================

def formatar_excel(
    caminho
):

    print(
        "\nFormatando arquivo Excel..."
    )

    wb = load_workbook(
        caminho
    )

    for ws in wb.worksheets:

        # ----------------------------------------------------
        # Congelar primeira linha
        # ----------------------------------------------------

        ws.freeze_panes = "A2"

        # ----------------------------------------------------
        # Filtros
        # ----------------------------------------------------

        if ws.max_row > 1:

            ws.auto_filter.ref = (
                ws.dimensions
            )

        # ----------------------------------------------------
        # Cabeçalho
        # ----------------------------------------------------

        for cell in ws[1]:

            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        # ----------------------------------------------------
        # Largura automática
        # ----------------------------------------------------

        for coluna in ws.columns:

            letra = (
                coluna[
                    0
                ].column_letter
            )

            maior = 0

            for cell in coluna:

                if cell.value is None:
                    continue

                tamanho = len(
                    str(cell.value)
                )

                if tamanho > maior:
                    maior = tamanho

            largura = min(
                max(
                    maior + 2,
                    10
                ),
                45
            )

            ws.column_dimensions[
                letra
            ].width = largura

        # ----------------------------------------------------
        # Formatação de datas
        # ----------------------------------------------------

        for row in ws.iter_rows():

            for cell in row:

                if isinstance(
                    cell.value,
                    datetime
                ):

                    cell.number_format = (
                        "yyyy-mm-dd hh:mm:ss"
                    )

    wb.save(
        caminho
    )


# ============================================================
# SALVAR RESULTADOS
# ============================================================

def salvar_resultados(
    eletricos,
    integracao_15,
    validacao,
    resumo
):

    print(
        "\nSalvando resultados..."
    )

    with pd.ExcelWriter(
        ARQUIVO_SAIDA,
        engine="openpyxl"
    ) as writer:

        eletricos.to_excel(
            writer,
            sheet_name="eletrico_5min",
            index=False
        )

        integracao_15.to_excel(
            writer,
            sheet_name="integracao_15min",
            index=False
        )

        validacao.to_excel(
            writer,
            sheet_name="validacao_potencia",
            index=False
        )

        resumo.to_excel(
            writer,
            sheet_name="resumo",
            index=False
        )

    formatar_excel(
        ARQUIVO_SAIDA
    )

    print(
        "\nArquivo final salvo em:"
    )

    print(
        os.path.abspath(
            ARQUIVO_SAIDA
        )
    )


# ============================================================
# MAIN
# ============================================================

def main():

    limpar_terminal()

    print(
        "=" * 70
    )

    print(
        "INTEGRAÇÃO METEOROLOGIA + USINA FV"
    )

    print(
        "22/09/2025"
    )

    print(
        "=" * 70
    )

    # --------------------------------------------------------
    # 1. VERIFICAR ARQUIVOS
    # --------------------------------------------------------

    verificar_arquivos()

    # --------------------------------------------------------
    # 2. USINA
    # --------------------------------------------------------

    usina = ler_relatorio_usina()

    # --------------------------------------------------------
    # 3. INVERSORES
    # --------------------------------------------------------

    inversores = (
        ler_relatorio_inversores()
    )

    # --------------------------------------------------------
    # 4. PREPARAR DADOS ELÉTRICOS
    # --------------------------------------------------------

    eletricos = (
        preparar_dados_eletricos(
            usina,
            inversores
        )
    )

    # --------------------------------------------------------
    # 5. METEOROLOGIA
    # --------------------------------------------------------

    meteorologia = (
        preparar_meteorologia()
    )

    # --------------------------------------------------------
    # 6. INTEGRAÇÃO EM 15 MINUTOS
    # --------------------------------------------------------

    integracao_15 = (
        integrar_15_minutos(
            eletricos,
            meteorologia
        )
    )

    # --------------------------------------------------------
    # 7. VALIDAÇÃO DA POTÊNCIA
    # --------------------------------------------------------

    (
        validacao,
        validacao_resumo
    ) = validar_potencia(
        eletricos
    )

    # --------------------------------------------------------
    # 8. RESUMO
    # --------------------------------------------------------

    resumo = gerar_resumo(
        eletricos,
        integracao_15,
        validacao_resumo
    )

    # --------------------------------------------------------
    # 9. SALVAR
    # --------------------------------------------------------

    salvar_resultados(
        eletricos,
        integracao_15,
        validacao,
        resumo
    )

    # --------------------------------------------------------
    # FINAL
    # --------------------------------------------------------

    print(
        "\n" + "=" * 70
    )

    print(
        "INTEGRAÇÃO CONCLUÍDA"
    )

    print(
        "=" * 70
    )

    print(
        "\nAbas criadas:"
    )

    print(
        "1. eletrico_5min"
    )

    print(
        "2. integracao_15min"
    )

    print(
        "3. validacao_potencia"
    )

    print(
        "4. resumo"
    )


# ============================================================
# EXECUÇÃO
# ============================================================

if __name__ == "__main__":

    main()