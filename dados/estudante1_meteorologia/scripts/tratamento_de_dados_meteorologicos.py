import os
import re
import numpy as np
import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment


# ============================================================
# CONFIGURAÇÕES
# ============================================================

# Descobre automaticamente onde o arquivo .py está.
# Assim o programa não depende da pasta em que o terminal
# foi aberto.

PASTA_PROJETO = os.path.dirname(
    os.path.abspath(__file__)
)

PASTA_DADOS = os.path.join(
    PASTA_PROJETO,
    "dados"
)

ARQUIVO_ENTRADA = os.path.join(
    PASTA_DADOS,
    "dados_meteorologicos.xlsx"
)

ARQUIVO_SAIDA = os.path.join(
    PASTA_DADOS,
    "dados_meteorologicos_tratados.xlsx"
)

# Frequência esperada da estação
INTERVALO_ESPERADO = pd.Timedelta(
    minutes=15
)

# Os timestamps possuem segundos que podem variar.
# Exemplo:
# 14:15:07 -> 14:30:08
# Isso não deve ser considerado uma lacuna.
TOLERANCIA_INTERVALO = pd.Timedelta(
    seconds=2
)


# ============================================================
# LIMPAR TERMINAL
# ============================================================

def limpar_terminal():

    if os.name == "nt":
        os.system("cls")
    else:
        os.system("clear")


# ============================================================
# NORMALIZAR NOME DE COLUNA
# ============================================================

def normalizar_nome(nome):

    nome = str(nome).strip().lower()

    substituicoes = {
        "á": "a",
        "à": "a",
        "ã": "a",
        "â": "a",
        "é": "e",
        "ê": "e",
        "í": "i",
        "ó": "o",
        "ô": "o",
        "õ": "o",
        "ú": "u",
        "ç": "c"
    }

    for antigo, novo in substituicoes.items():
        nome = nome.replace(
            antigo,
            novo
        )

    return nome


# ============================================================
# ENCONTRAR COLUNA
# ============================================================

def encontrar_coluna(df, nome):

    alvo = normalizar_nome(nome)

    for coluna in df.columns:

        if normalizar_nome(coluna) == alvo:
            return coluna

    return None


# ============================================================
# CARREGAR DADOS
# ============================================================

def carregar_dados():

    print("\n" + "=" * 70)
    print("TRATAMENTO COMPLETO DOS DADOS METEOROLÓGICOS")
    print("=" * 70)

    print("\nCarregando arquivo...")

    if not os.path.exists(ARQUIVO_ENTRADA):

        print("\nERRO: arquivo não encontrado.")
        print(
            f"\nCaminho procurado:\n{ARQUIVO_ENTRADA}"
        )

        raise FileNotFoundError(
            ARQUIVO_ENTRADA
        )

    df = pd.read_excel(
        ARQUIVO_ENTRADA
    )

    print("Arquivo carregado com sucesso.")
    print(
        f"Linhas encontradas: {len(df)}"
    )
    print(
        f"Colunas encontradas: {len(df.columns)}"
    )

    # Remove somente linhas completamente vazias.
    antes = len(df)

    df = df.dropna(
        how="all"
    ).copy()

    removidas = antes - len(df)

    print(
        f"\nLinhas completamente vazias removidas: "
        f"{removidas}"
    )

    return df


# ============================================================
# PREPARAR TIMESTAMPS
# ============================================================

def preparar_timestamps(df):

    print("\nPreparando timestamps...")

    coluna_tempo = None

    # Preferimos UTC -3 porque é o horário local.
    for nome in [
        "UTC -3",
        "Local",
        "UTC"
    ]:

        coluna = encontrar_coluna(
            df,
            nome
        )

        if coluna is not None:

            coluna_tempo = coluna
            break

    if coluna_tempo is None:

        raise ValueError(
            "Nenhuma coluna de timestamp encontrada."
        )

    print(
        f"Coluna temporal utilizada: "
        f"{coluna_tempo}"
    )

    df[coluna_tempo] = pd.to_datetime(
        df[coluna_tempo],
        errors="coerce"
    )

    invalidos = df[
        coluna_tempo
    ].isna().sum()

    print(
        f"Timestamps inválidos ou ausentes: "
        f"{invalidos}"
    )

    if invalidos > 0:

        df = df.loc[
            df[coluna_tempo].notna()
        ].copy()

    # Mantém uma coluna padronizada chamada UTC -3.
    if coluna_tempo != "UTC -3":

        df["UTC -3"] = pd.to_datetime(
            df[coluna_tempo],
            errors="coerce"
        )

    print(
        f"Registros removidos por timestamp inválido: "
        f"{invalidos}"
    )

    return df


# ============================================================
# ORDENAR CRONOLOGICAMENTE
# ============================================================

def ordenar_registros(df):

    print(
        "\nOrdenando registros cronologicamente..."
    )

    df = df.sort_values(
        by="UTC -3"
    ).reset_index(
        drop=True
    )

    print(
        "Registros ordenados."
    )

    return df


# ============================================================
# CONVERTER COLUNAS NUMÉRICAS
# ============================================================

def converter_colunas_numericas(df):

    print(
        "\nVerificando colunas numéricas..."
    )

    ignorar = {
        "UTC",
        "UTC -3",
        "Local",
        "Freq"
    }

    convertidas = []

    for coluna in df.columns:

        if coluna in ignorar:
            continue

        if pd.api.types.is_numeric_dtype(
            df[coluna]
        ):
            continue

        original = df[coluna]

        quantidade_original = (
            original.notna().sum()
        )

        if quantidade_original == 0:
            continue

        convertido = pd.to_numeric(
            original,
            errors="coerce"
        )

        quantidade_convertida = (
            convertido.notna().sum()
        )

        proporcao = (
            quantidade_convertida
            / quantidade_original
        )

        # Só converte quando pelo menos 95%
        # dos valores podem ser interpretados
        # como números.
        if proporcao >= 0.95:

            df[coluna] = convertido

            convertidas.append(
                coluna
            )

    print(
        f"Colunas convertidas para numérico: "
        f"{len(convertidas)}"
    )

    return df


# ============================================================
# DUPLICADOS
# ============================================================

def analisar_duplicados(df):

    print(
        "\nAnalisando timestamps duplicados..."
    )

    mascara = df[
        "UTC -3"
    ].duplicated(
        keep=False
    )

    quantidade = mascara.sum()

    print(
        f"Registros com timestamp duplicado: "
        f"{quantidade}"
    )

    if quantidade > 0:

        duplicados = df.loc[
            mascara
        ].copy()

    else:

        duplicados = pd.DataFrame(
            columns=df.columns
        )

    return duplicados


# ============================================================
# INTERVALOS TEMPORAIS
# ============================================================

def analisar_intervalos(df):

    print(
        "\nAnalisando intervalos temporais..."
    )

    df["intervalo_anterior"] = (
        df["UTC -3"].diff()
    )

    # Intervalos negativos.
    df["intervalo_negativo"] = (
        df["intervalo_anterior"]
        < pd.Timedelta(0)
    )

    limite_lacuna = (
        INTERVALO_ESPERADO
        + TOLERANCIA_INTERVALO
    )

    df["lacuna_temporal"] = (
        df["intervalo_anterior"]
        > limite_lacuna
    )

    quantidade_lacunas = int(
        df["lacuna_temporal"].sum()
    )

    quantidade_negativos = int(
        df["intervalo_negativo"].sum()
    )

    print(
        f"Lacunas temporais: "
        f"{quantidade_lacunas}"
    )

    print(
        f"Intervalos negativos: "
        f"{quantidade_negativos}"
    )

    print(
        f"\nTotal de lacunas encontradas: "
        f"{quantidade_lacunas}"
    )

    return df


# ============================================================
# GERAR TABELA DE LACUNAS
# ============================================================

def gerar_tabela_lacunas(df):

    lacunas = df.loc[
        df["lacuna_temporal"],
        [
            "UTC -3",
            "intervalo_anterior"
        ]
    ].copy()

    if lacunas.empty:

        return pd.DataFrame(
            columns=[
                "ultimo_registro_anterior",
                "primeiro_registro_posterior",
                "duracao_horas",
                "registros_ausentes"
            ]
        )

    # Primeiro registro depois da lacuna.
    lacunas[
        "primeiro_registro_posterior"
    ] = lacunas[
        "UTC -3"
    ]

    # Último registro antes da lacuna.
    lacunas[
        "ultimo_registro_anterior"
    ] = (
        lacunas[
            "primeiro_registro_posterior"
        ]
        - lacunas[
            "intervalo_anterior"
        ]
    )

    # Duração em horas.
    lacunas[
        "duracao_horas"
    ] = (
        lacunas[
            "intervalo_anterior"
        ]
        / pd.Timedelta(
            hours=1
        )
    )

    # Quantidade de registros de 15 minutos
    # que deveriam existir dentro do intervalo.
    #
    # Exemplo:
    #
    # 10:00
    # 10:15
    # 10:30
    #
    # Se só houver 10:00 e 10:30:
    #
    # 30 / 15 = 2 intervalos
    # 2 - 1 = 1 registro ausente.

    lacunas[
        "registros_ausentes"
    ] = (
        (
            lacunas[
                "intervalo_anterior"
            ]
            / INTERVALO_ESPERADO
        )
        .round()
        - 1
    )

    lacunas[
        "registros_ausentes"
    ] = (
        lacunas[
            "registros_ausentes"
        ]
        .clip(lower=1)
        .astype(int)
    )

    lacunas = lacunas[
        [
            "ultimo_registro_anterior",
            "primeiro_registro_posterior",
            "duracao_horas",
            "registros_ausentes"
        ]
    ]

    return lacunas.reset_index(
        drop=True
    )


# ============================================================
# ENCONTRAR GRUPOS MIN / MED / MAX
# ============================================================

def encontrar_grupos_min_med_max(df):

    grupos = {}

    for coluna in df.columns:

        nome = normalizar_nome(
            coluna
        )

        correspondencia = re.search(
            r"(?:^|[\s_])(min|med|max)$",
            nome
        )

        if correspondencia is None:
            continue

        sufixo = correspondencia.group(1)

        # Remove min/med/max do final.
        base = re.sub(
            r"[\s_]+(?:min|med|max)$",
            "",
            str(coluna),
            flags=re.IGNORECASE
        ).strip()

        chave = normalizar_nome(
            base
        )

        if chave not in grupos:

            grupos[chave] = {
                "base": base,
                "min": None,
                "med": None,
                "max": None
            }

        grupos[chave][
            sufixo
        ] = coluna

    resultado = []

    for grupo in grupos.values():

        if (
            grupo["min"] is not None
            and grupo["med"] is not None
            and grupo["max"] is not None
        ):

            resultado.append(
                (
                    grupo["base"],
                    grupo["min"],
                    grupo["med"],
                    grupo["max"]
                )
            )

    return resultado


# ============================================================
# VALIDAR MIN / MED / MAX
# ============================================================

def validar_min_med_max(df):

    print(
        "\nVerificando relações "
        "mínimo <= média <= máximo..."
    )

    grupos = encontrar_grupos_min_med_max(
        df
    )

    resultados = []

    for (
        base,
        coluna_min,
        coluna_med,
        coluna_max
    ) in grupos:

        # Direção do vento NÃO recebe esta validação,
        # porque graus são uma variável circular.
        if (
            "direcao vento"
            in normalizar_nome(base)
        ):
            continue

        minimo = pd.to_numeric(
            df[coluna_min],
            errors="coerce"
        )

        media = pd.to_numeric(
            df[coluna_med],
            errors="coerce"
        )

        maximo = pd.to_numeric(
            df[coluna_max],
            errors="coerce"
        )

        mascara = (
            (minimo > media)
            |
            (media > maximo)
        )

        quantidade = int(
            mascara.sum()
        )

        resultados.append(
            {
                "variavel": base,
                "coluna_min": coluna_min,
                "coluna_med": coluna_med,
                "coluna_max": coluna_max,
                "inconsistencias": quantidade
            }
        )

    resultado = pd.DataFrame(
        resultados
    )

    print(
        f"Grupos min/med/max analisados: "
        f"{len(resultado)}"
    )

    if not resultado.empty:

        print(
            f"Inconsistências encontradas: "
            f"{resultado['inconsistencias'].sum()}"
        )

    else:

        print(
            "Nenhum grupo min/med/max encontrado."
        )

    return resultado


# ============================================================
# CRIAR FLAGS DE QUALIDADE
# ============================================================

def criar_flags_fisicas(df):

    print(
        "\nCriando flags de qualidade física..."
    )

    flags = {}

    def adicionar_flag(
        coluna,
        mascara,
        descricao
    ):

        if coluna not in df.columns:
            return

        flags[coluna] = {
            "mascara": mascara.fillna(
                False
            ).astype(bool),
            "descricao": descricao
        }

    # --------------------------------------------------------
    # TEMPERATURA
    # --------------------------------------------------------

    coluna = encontrar_coluna(
        df,
        "Temperatura inst"
    )

    if coluna is not None:

        serie = pd.to_numeric(
            df[coluna],
            errors="coerce"
        )

        mascara = (
            serie == -40
        )

        adicionar_flag(
            coluna,
            mascara,
            "Valor -40 observado como código/valor inválido."
        )

    # --------------------------------------------------------
    # UMIDADE RELATIVA
    # --------------------------------------------------------

    for coluna in df.columns:

        if (
            "umidade relativa"
            in normalizar_nome(coluna)
        ):

            serie = pd.to_numeric(
                df[coluna],
                errors="coerce"
            )

            mascara = (
                (serie < 0)
                |
                (serie > 100)
            )

            adicionar_flag(
                coluna,
                mascara,
                "Umidade relativa fora de 0% a 100%."
            )

    # --------------------------------------------------------
    # DIREÇÃO DO VENTO
    # --------------------------------------------------------

    for coluna in df.columns:

        if (
            "direcao vento"
            in normalizar_nome(coluna)
        ):

            serie = pd.to_numeric(
                df[coluna],
                errors="coerce"
            )

            mascara = (
                (serie < 0)
                |
                (serie > 360)
            )

            adicionar_flag(
                coluna,
                mascara,
                "Direção do vento fora de 0° a 360°."
            )

    # --------------------------------------------------------
    # VELOCIDADE DO VENTO
    # --------------------------------------------------------

    for coluna in df.columns:

        if (
            "velocidade vento"
            in normalizar_nome(coluna)
        ):

            serie = pd.to_numeric(
                df[coluna],
                errors="coerce"
            )

            mascara = (
                serie < 0
            )

            adicionar_flag(
                coluna,
                mascara,
                "Velocidade do vento negativa."
            )

    # --------------------------------------------------------
    # PRECIPITAÇÃO
    # --------------------------------------------------------

    coluna = encontrar_coluna(
        df,
        "Precipitação"
    )

    if coluna is not None:

        serie = pd.to_numeric(
            df[coluna],
            errors="coerce"
        )

        mascara = (
            serie < 0
        )

        adicionar_flag(
            coluna,
            mascara,
            "Precipitação negativa."
        )

    # --------------------------------------------------------
    # RADIAÇÃO SOLAR GLOBAL
    # --------------------------------------------------------

    for coluna in df.columns:

        if (
            "radiacao solar global"
            in normalizar_nome(coluna)
        ):

            serie = pd.to_numeric(
                df[coluna],
                errors="coerce"
            )

            mascara = (
                serie < 0
            )

            adicionar_flag(
                coluna,
                mascara,
                "Radiação solar global negativa."
            )

    # --------------------------------------------------------
    # UMIDADE DO SOLO
    # --------------------------------------------------------

    for coluna in df.columns:

        if (
            "umid solo"
            in normalizar_nome(coluna)
        ):

            serie = pd.to_numeric(
                df[coluna],
                errors="coerce"
            )

            mascara = (
                serie == -400
            )

            adicionar_flag(
                coluna,
                mascara,
                "Valor -400 identificado como possível código de erro."
            )

    print(
        f"Variáveis com regras de validação: "
        f"{len(flags)}"
    )

    return flags


# ============================================================
# APLICAR FLAGS
# ============================================================

def aplicar_flags(df, flags):

    print(
        "\nAplicando flags ao conjunto de dados..."
    )

    for coluna, dados in flags.items():

        nome = normalizar_nome(
            coluna
        )

        nome = re.sub(
            r"[^a-zA-Z0-9]+",
            "_",
            nome
        ).strip("_")

        nome_flag = (
            f"flag_{nome}"
        )

        df[nome_flag] = (
            dados["mascara"]
            .astype(bool)
        )

    return df


# ============================================================
# CRIAR VALORES TRATADOS
# ============================================================

def criar_dados_tratados(
    df,
    flags
):

    print(
        "\nCriando valores tratados..."
    )

    tratado = df.copy()

    for coluna, dados in flags.items():

        mascara = dados[
            "mascara"
        ]

        # Converte para Float64, que permite
        # valores numéricos e NA.
        serie = pd.to_numeric(
            tratado[coluna],
            errors="coerce"
        ).astype(
            "Float64"
        )

        serie.loc[
            mascara
        ] = pd.NA

        tratado[
            f"{coluna}_tratado"
        ] = serie

    return tratado


# ============================================================
# ANÁLISE DE VALORES SUSPEITOS
# ============================================================

def analisar_valores_suspeitos(df):

    print(
        "\nGerando análise de valores suspeitos..."
    )

    resultados = []

    def adicionar(
        grupo,
        tipo,
        quantidade,
        colunas,
        observacao
    ):

        resultados.append(
            {
                "variavel_grupo": grupo,
                "tipo": tipo,
                "quantidade": int(quantidade),
                "colunas_analisadas": ", ".join(colunas),
                "observacao": observacao
            }
        )

    # ========================================================
    # TEMPERATURA INST = -40
    # ========================================================

    coluna = encontrar_coluna(
        df,
        "Temperatura inst"
    )

    if coluna is not None:

        serie = pd.to_numeric(
            df[coluna],
            errors="coerce"
        )

        quantidade = int(
            (serie == -40).sum()
        )

        adicionar(
            "Temperatura",
            "INVÁLIDO",
            quantidade,
            [coluna],
            "Valor -40 identificado como código/valor inválido."
        )

    # ========================================================
    # RADIAÇÃO IV
    # ========================================================

    colunas = [
        coluna
        for coluna in df.columns
        if "radiacao iv"
        in normalizar_nome(coluna)
    ]

    if colunas:

        total_negativos = 0

        for coluna in colunas:

            serie = pd.to_numeric(
                df[coluna],
                errors="coerce"
            )

            total_negativos += int(
                (serie < 0).sum()
            )

        adicionar(
            "Radiação IV",
            "SUSPEITO",
            total_negativos,
            colunas,
            "Valores negativos encontrados nas séries de Radiação IV. Mantidos para análise até confirmação da unidade, escala e convenção do sensor."
        )

    # ========================================================
    # UMIDADE DO SOLO
    # ========================================================

    colunas = [
        coluna
        for coluna in df.columns
        if "umid solo"
        in normalizar_nome(coluna)
    ]

    if colunas:

        total_negativos = 0

        for coluna in colunas:

            serie = pd.to_numeric(
                df[coluna],
                errors="coerce"
            )

            total_negativos += int(
                (serie < 0).sum()
            )

        adicionar(
            "Umidade do solo",
            "SUSPEITO",
            total_negativos,
            colunas,
            "Valores negativos encontrados nas séries de umidade do solo. Mantidos para análise até confirmação da unidade, escala e convenção do sensor."
        )

    # ========================================================
    # DIREÇÃO DO VENTO
    # ========================================================

    colunas = [
        coluna
        for coluna in df.columns
        if "direcao vento"
        in normalizar_nome(coluna)
    ]

    if colunas:

        total_invalidos = 0

        for coluna in colunas:

            serie = pd.to_numeric(
                df[coluna],
                errors="coerce"
            )

            total_invalidos += int(
                (
                    (serie < 0)
                    |
                    (serie > 360)
                ).sum()
            )

        adicionar(
            "Direção do vento",
            "INVÁLIDO",
            total_invalidos,
            colunas,
            "Direção fora do intervalo de 0° a 360°."
        )

    # ========================================================
    # VELOCIDADE DO VENTO
    # ========================================================

    colunas = [
        coluna
        for coluna in df.columns
        if "velocidade vento"
        in normalizar_nome(coluna)
    ]

    if colunas:

        total_invalidos = 0

        for coluna in colunas:

            serie = pd.to_numeric(
                df[coluna],
                errors="coerce"
            )

            total_invalidos += int(
                (serie < 0).sum()
            )

        adicionar(
            "Velocidade do vento",
            "INVÁLIDO",
            total_invalidos,
            colunas,
            "Velocidade do vento negativa."
        )

    # ========================================================
    # UMIDADE RELATIVA
    # ========================================================

    colunas = [
        coluna
        for coluna in df.columns
        if "umidade relativa"
        in normalizar_nome(coluna)
    ]

    if colunas:

        total_invalidos = 0

        for coluna in colunas:

            serie = pd.to_numeric(
                df[coluna],
                errors="coerce"
            )

            total_invalidos += int(
                (
                    (serie < 0)
                    |
                    (serie > 100)
                ).sum()
            )

        adicionar(
            "Umidade relativa",
            "INVÁLIDO",
            total_invalidos,
            colunas,
            "Umidade relativa fora do intervalo de 0% a 100%."
        )

    # ========================================================
    # PRECIPITAÇÃO
    # ========================================================

    coluna = encontrar_coluna(
        df,
        "Precipitação"
    )

    if coluna is not None:

        serie = pd.to_numeric(
            df[coluna],
            errors="coerce"
        )

        quantidade = int(
            (serie < 0).sum()
        )

        adicionar(
            "Precipitação",
            "INVÁLIDO",
            quantidade,
            [coluna],
            "Precipitação negativa."
        )

    # ========================================================
    # RADIAÇÃO SOLAR GLOBAL
    # ========================================================

    colunas = [
        coluna
        for coluna in df.columns
        if "radiacao solar global"
        in normalizar_nome(coluna)
    ]

    if colunas:

        total_invalidos = 0

        for coluna in colunas:

            serie = pd.to_numeric(
                df[coluna],
                errors="coerce"
            )

            total_invalidos += int(
                (serie < 0).sum()
            )

        adicionar(
            "Radiação solar global",
            "INVÁLIDO",
            total_invalidos,
            colunas,
            "Radiação solar global negativa."
        )

    return pd.DataFrame(
        resultados
    )


# ============================================================
# COLUNAS DE MEDIÇÃO
# ============================================================

def obter_colunas_medicao(df):

    ignorar = {
        "UTC",
        "UTC -3",
        "Local",
        "Freq",
        "intervalo_anterior",
        "intervalo_negativo",
        "lacuna_temporal"
    }

    colunas = []

    for coluna in df.columns:

        if coluna in ignorar:
            continue

        if str(coluna).startswith(
            "flag_"
        ):
            continue

        if str(coluna).endswith(
            "_tratado"
        ):
            continue

        if pd.api.types.is_numeric_dtype(
            df[coluna]
        ):

            colunas.append(
                coluna
            )

    return colunas


# ============================================================
# RELATÓRIO DAS VARIÁVEIS
# ============================================================

def gerar_relatorio_variaveis(
    df_original,
    df_tratado,
    flags
):

    print(
        "\nGerando relatório das variáveis..."
    )

    colunas = obter_colunas_medicao(
        df_original
    )

    resultados = []

    for coluna in colunas:

        original = pd.to_numeric(
            df_original[coluna],
            errors="coerce"
        )

        coluna_tratada = (
            f"{coluna}_tratado"
        )

        if coluna_tratada in df_tratado.columns:

            tratado = pd.to_numeric(
                df_tratado[
                    coluna_tratada
                ],
                errors="coerce"
            )

        else:

            tratado = original.copy()

        ausentes_original = int(
            original.isna().sum()
        )

        ausentes_tratado = int(
            tratado.isna().sum()
        )

        if coluna in flags:

            invalidados = int(
                flags[coluna][
                    "mascara"
                ].sum()
            )

        else:

            invalidados = 0

        validos = tratado.dropna()

        if len(validos) > 0:

            minimo = validos.min()
            maximo = validos.max()
            media = validos.mean()

        else:

            minimo = np.nan
            maximo = np.nan
            media = np.nan

        resultados.append(
            {
                "variavel": coluna,
                "registros": len(
                    df_original
                ),
                "ausentes_original": ausentes_original,
                "ausentes_tratado": ausentes_tratado,
                "valores_invalidados": invalidados,
                "minimo": minimo,
                "maximo": maximo,
                "media": media
            }
        )

    return pd.DataFrame(
        resultados
    )


# ============================================================
# RESUMO
# ============================================================

def gerar_resumo(
    df,
    lacunas,
    duplicados,
    min_med_max,
    relatorio,
    suspeitos
):

    print(
        "\nGerando resumo geral..."
    )

    primeiro = df[
        "UTC -3"
    ].min()

    ultimo = df[
        "UTC -3"
    ].max()

    if not lacunas.empty:

        registros_ausentes = int(
            lacunas[
                "registros_ausentes"
            ].sum()
        )

    else:

        registros_ausentes = 0

    if not min_med_max.empty:

        inconsistencias = int(
            min_med_max[
                "inconsistencias"
            ].sum()
        )

    else:

        inconsistencias = 0

    valores_invalidados = int(
        relatorio[
            "valores_invalidados"
        ].sum()
    )

    resumo = pd.DataFrame(
        {
            "indicador": [
                "Total de registros",
                "Total de colunas",
                "Primeiro registro",
                "Último registro",
                "Intervalo esperado",
                "Tolerância temporal",
                "Timestamps duplicados",
                "Lacunas temporais",
                "Registros ausentes estimados",
                "Valores invalidados",
                "Inconsistências min/med/max",
                "Variáveis analisadas",
                "Ocorrências suspeitas"
            ],
            "valor": [
                len(df),
                len(df.columns),
                primeiro,
                ultimo,
                "15 minutos",
                "2 segundos",
                len(duplicados),
                len(lacunas),
                registros_ausentes,
                valores_invalidados,
                inconsistencias,
                len(relatorio),
                len(suspeitos)
            ]
        }
    )

    return resumo


# ============================================================
# FORMATAÇÃO DO EXCEL
# ============================================================

def formatar_excel(caminho):

    print(
        "\nFormatando arquivo Excel..."
    )

    wb = load_workbook(
        caminho
    )

    for ws in wb.worksheets:

        # Congela a primeira linha.
        ws.freeze_panes = "A2"

        # Filtro automático.
        if ws.max_row > 1:
            ws.auto_filter.ref = (
                ws.dimensions
            )

        # Cabeçalho.
        for cell in ws[1]:

            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        # Ajusta automaticamente as colunas.
        for coluna in ws.columns:

            letra = (
                coluna[0]
                .column_letter
            )

            maior = 0

            for cell in coluna:

                if cell.value is None:
                    continue

                tamanho = len(
                    str(cell.value)
                )

                if tamanho > maior:
                    maior = tamanho

            largura = min(
                max(
                    maior + 2,
                    10
                ),
                45
            )

            ws.column_dimensions[
                letra
            ].width = largura

        # Formatação de timestamps.
        for row in ws.iter_rows():

            for cell in row:

                if isinstance(
                    cell.value,
                    (
                        pd.Timestamp,
                    )
                ):

                    cell.number_format = (
                        "yyyy-mm-dd hh:mm:ss"
                    )

    wb.save(
        caminho
    )


# ============================================================
# SALVAR RESULTADOS
# ============================================================

def salvar_resultados(
    df_tratado,
    lacunas,
    duplicados,
    min_med_max,
    relatorio,
    suspeitos,
    resumo
):

    print(
        "\nSalvando resultados..."
    )

    with pd.ExcelWriter(
        ARQUIVO_SAIDA,
        engine="openpyxl"
    ) as writer:

        df_tratado.to_excel(
            writer,
            sheet_name="dados_tratados",
            index=False
        )

        lacunas.to_excel(
            writer,
            sheet_name="lacunas",
            index=False
        )

        duplicados.to_excel(
            writer,
            sheet_name="duplicados",
            index=False
        )

        min_med_max.to_excel(
            writer,
            sheet_name="min_med_max",
            index=False
        )

        relatorio.to_excel(
            writer,
            sheet_name="qualidade_variaveis",
            index=False
        )

        suspeitos.to_excel(
            writer,
            sheet_name="suspeitos",
            index=False
        )

        resumo.to_excel(
            writer,
            sheet_name="resumo",
            index=False
        )

    formatar_excel(
        ARQUIVO_SAIDA
    )

    print(
        "\nArquivo final salvo em:"
    )

    print(
        os.path.abspath(
            ARQUIVO_SAIDA
        )
    )


# ============================================================
# MAIN
# ============================================================

def main():

    # Limpa o terminal antes de começar.
    limpar_terminal()

    # --------------------------------------------------------
    # 1. CARREGAR
    # --------------------------------------------------------

    df = carregar_dados()

    # --------------------------------------------------------
    # 2. TIMESTAMPS
    # --------------------------------------------------------

    df = preparar_timestamps(
        df
    )

    # --------------------------------------------------------
    # 3. ORDENAR
    # --------------------------------------------------------

    df = ordenar_registros(
        df
    )

    # --------------------------------------------------------
    # 4. CONVERTER NÚMEROS
    # --------------------------------------------------------

    df = converter_colunas_numericas(
        df
    )

    # --------------------------------------------------------
    # 5. DUPLICADOS
    # --------------------------------------------------------

    duplicados = analisar_duplicados(
        df
    )

    # --------------------------------------------------------
    # 6. INTERVALOS
    # --------------------------------------------------------

    df = analisar_intervalos(
        df
    )

    # --------------------------------------------------------
    # 7. TABELA DE LACUNAS
    # --------------------------------------------------------

    lacunas = gerar_tabela_lacunas(
        df
    )

    # --------------------------------------------------------
    # 8. MIN / MED / MAX
    # --------------------------------------------------------

    min_med_max = validar_min_med_max(
        df
    )

    # --------------------------------------------------------
    # 9. FLAGS FÍSICAS
    # --------------------------------------------------------

    flags = criar_flags_fisicas(
        df
    )

    # --------------------------------------------------------
    # 10. APLICAR FLAGS
    # --------------------------------------------------------

    df_com_flags = aplicar_flags(
        df.copy(),
        flags
    )

    # --------------------------------------------------------
    # 11. CRIAR VALORES TRATADOS
    # --------------------------------------------------------

    df_tratado = criar_dados_tratados(
        df_com_flags,
        flags
    )

    # --------------------------------------------------------
    # 12. SUSPEITOS
    # --------------------------------------------------------

    suspeitos = analisar_valores_suspeitos(
        df
    )

    # --------------------------------------------------------
    # 13. RELATÓRIO DAS VARIÁVEIS
    # --------------------------------------------------------

    relatorio = gerar_relatorio_variaveis(
        df,
        df_tratado,
        flags
    )

    # --------------------------------------------------------
    # 14. RESUMO
    # --------------------------------------------------------

    resumo = gerar_resumo(
        df,
        lacunas,
        duplicados,
        min_med_max,
        relatorio,
        suspeitos
    )

    # --------------------------------------------------------
    # 15. SALVAR
    # --------------------------------------------------------

    salvar_resultados(
        df_tratado,
        lacunas,
        duplicados,
        min_med_max,
        relatorio,
        suspeitos,
        resumo
    )

    # --------------------------------------------------------
    # FINAL
    # --------------------------------------------------------

    print("\n" + "=" * 70)
    print("TRATAMENTO CONCLUÍDO")
    print("=" * 70)

    print("\nAbas criadas no arquivo:")

    print("1. dados_tratados")
    print("2. lacunas")
    print("3. duplicados")
    print("4. min_med_max")
    print("5. qualidade_variaveis")
    print("6. suspeitos")
    print("7. resumo")


# ============================================================
# EXECUÇÃO
# ============================================================

if __name__ == "__main__":
    main()